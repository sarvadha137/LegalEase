from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import PlainTextResponse
from pathlib import Path
from fastapi.middleware.cors import CORSMiddleware
from fpdf import FPDF
from fastapi.responses import Response
import os
from dotenv import load_dotenv
from google import genai
from docx import Document

app = FastAPI(title="LegalEase API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

load_dotenv()

client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

from backend.document_service import (
    extract_text_from_pdf,
    extract_text_from_docx
)


UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


@app.get("/")
def home():
    return {
        "message": "Welcome to LegalEase!",
        "status": "Backend is running"
    }


@app.get("/health")
def health_check():
    return {
        "status": "healthy"
    }


@app.post("/upload-document")
async def upload_document(file: UploadFile = File(...)):

    allowed_extensions = [".pdf", ".docx"]

    file_extension = Path(file.filename).suffix.lower()

    if file_extension not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are allowed."
        )

    file_path = UPLOAD_DIR / file.filename

    content = await file.read()

    with open(file_path, "wb") as buffer:
        buffer.write(content)

    if file_extension == ".pdf":
        text = extract_text_from_pdf(file_path)

    else:
        text = extract_text_from_docx(file_path)

    return {
        "filename": file.filename,
        "message": "Document uploaded successfully",
        "text_preview": text[:1000]
    }
@app.post("/summarize-text")
def summarize_text(text: str):

    prompt = f"""
You are a legal document assistant.

Analyze the following legal document and provide a clear, structured summary.

Include:
1. Document type
2. Parties involved
3. Important dates
4. Financial terms
5. Main obligations
6. Notice or termination terms
7. Important clauses
8. Potential points the user should pay attention to

Use simple language.
Do not provide definitive legal advice.

Legal document:

{text}
"""

    response = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=prompt
    )

    return {
        "summary": response.text
    }
from pydantic import BaseModel


class DocumentRequest(BaseModel):
    document_type: str
    parties: str
    terms: str
    effective_date: str


@app.post("/generate-document")
def generate_document(request: DocumentRequest):

    prompt = f"""
You are a professional legal document drafting assistant.

Generate a clear, professionally structured legal document based on the
information provided below.

Document Type:
{request.document_type}

Parties:
{request.parties}

Terms and Conditions:
{request.terms}

Effective Date:
{request.effective_date}

Requirements:
- Use clear legal language.
- Include appropriate headings and sections.
- Include the parties and effective date.
- Include the provided terms and conditions.
- Make the document professionally structured.
- Do not invent important personal or financial information.
- If important information is missing, use a clear placeholder such as [INSERT INFORMATION].
- Include a short disclaimer that the document should be reviewed by a qualified legal professional.

Return only the drafted legal document.
"""

    response = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=prompt
    )

    return {
        "document": response.text
    }
   


class DownloadRequest(BaseModel):
    document: str


@app.post("/download-txt")
def download_txt(request: DownloadRequest):

    return PlainTextResponse(
        request.document,
        media_type="text/plain",
        headers={
            "Content-Disposition": "attachment; filename=legal_document.txt"
        }
    )
@app.post("/download-pdf")
def download_pdf(request: DownloadRequest):

    pdf = FPDF()
    pdf.add_page()

    font_path = Path("backend/fonts/arial.ttf")

    pdf.add_font(
        "ArialUnicode",
        "",
        str(font_path),
        uni=True
    )

    pdf.set_font("ArialUnicode", size=12)

    for line in request.document.split("\n"):
        safe_line = line.replace("₹", "INR ")
        pdf.multi_cell(0, 8, safe_line)

    pdf_output = pdf.output(dest="S")

    pdf_bytes = pdf_output.encode("latin-1")

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
            "attachment; filename=legal_document.pdf"
        }
    )
@app.post("/download-docx")
def download_docx(request: DownloadRequest):

    document = Document()

    # Title
    title = document.add_heading("LEGAL DOCUMENT", level=0)
    title.alignment = 1

    # Document content
    for line in request.document.split("\n"):
        if line.strip():
            document.add_paragraph(line.strip())

    # Save temporarily
    file_path = Path("uploads/legal_document.docx")
    document.save(file_path)

    with open(file_path, "rb") as file:
        docx_bytes = file.read()

    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition":
            "attachment; filename=legal_document.docx"
        }
    )